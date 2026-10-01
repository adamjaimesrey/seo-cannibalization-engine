"""Write this cycle's rows into the shared edit-history workbooks (Step 9 seed).

Masters live in EDIT_HISTORY_DIR (default data/edit-history).
Default is a DRY RUN. Pass --apply to write (a timestamped backup is made first).

Rules:
- Rows are matched by edit_id.
- An existing row is only filled if it is a PLACEHOLDER (its URL cell is empty).
- A real existing record (URL filled) is NEVER overwritten: it is skipped with a warning.
- edit_ids not present in the sheet are appended.
"""
import json
import os
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

MASTER_DIR = Path(os.environ.get("EDIT_HISTORY_DIR", "data/edit-history"))
BLOG_HISTORY = MASTER_DIR / "blog-edit-history.xlsx"
META_HISTORY = MASTER_DIR / "meta-edit-history.xlsx"
ROWS_FILE = Path("reference/history-rows-YYYY-MM.json")


def find_header(ws):
    for r in range(1, 11):
        values = [c.value for c in ws[r]]
        if "edit_id" in values:
            return r, {v: i + 1 for i, v in enumerate(values) if v}
    raise ValueError(f"No 'edit_id' header found in sheet '{ws.title}'")


def find_sheet(wb):
    for ws in wb.worksheets:
        try:
            find_header(ws)
            return ws
        except ValueError:
            continue
    raise ValueError("No sheet with an 'edit_id' header")


def last_data_row(ws, header_row, id_col):
    last = header_row
    for r in range(header_row + 1, ws.max_row + 1):
        if ws.cell(r, id_col).value not in (None, ""):
            last = r
    return last


def apply_rows(path, rows, url_field, apply):
    wb = load_workbook(path)
    ws = find_sheet(wb)
    header_row, cols = find_header(ws)
    missing = [k for k in rows[0] if k not in cols]
    if missing:
        raise ValueError(f"{path.name}: columns not found in sheet: {missing}")

    id_col = cols["edit_id"]
    index = {}
    for r in range(header_row + 1, ws.max_row + 1):
        v = ws.cell(r, id_col).value
        if v not in (None, ""):
            index[str(v).strip()] = r

    print(f"\n== {path.name} (sheet '{ws.title}', header row {header_row})")
    append_at = last_data_row(ws, header_row, id_col) + 1
    for row in rows:
        eid = row["edit_id"]
        if eid in index:
            r = index[eid]
            if ws.cell(r, cols[url_field]).value not in (None, ""):
                print(f"  SKIP  {eid}: real record already exists at row {r} (never overwritten)")
                continue
            action = f"FILL  {eid}: placeholder row {r}"
        else:
            r = append_at
            append_at += 1
            action = f"ADD   {eid}: new row {r}"
        print(f"  {action} -> {row[url_field]}")
        if apply:
            for field, value in row.items():
                ws.cell(r, cols[field]).value = value if value != "" else None

    if apply:
        backup = path.with_name(f"{path.stem}.backup-{datetime.now():%Y%m%d-%H%M%S}{path.suffix}")
        shutil.copy2(path, backup)
        wb.save(path)
        print(f"  saved. backup: {backup.name}")
    else:
        print("  (dry run, nothing written)")


def main():
    apply = "--apply" in sys.argv
    rows_file = Path(sys.argv[sys.argv.index("--rows") + 1]) if "--rows" in sys.argv else ROWS_FILE
    data = json.loads(rows_file.read_text(encoding="utf-8"))
    apply_rows(BLOG_HISTORY, data["blog_edit_history"], "blog_url", apply)
    apply_rows(META_HISTORY, data["meta_edit_history"], "page_url", apply)
    if not apply:
        print("\nDry run only. Re-run with --apply to write.")



def print_sheet_reminder():
    print("\nHITL: refresh the Google Sheet copies from the masters")
    print("  [ ] BEFORE importing: anything typed into the Google Sheet since the last import will be overwritten.")
    print("      Make sure it is also in the master xlsx.")
    print("  [ ] Open Google Sheet 'blog-edit-history', tab 'Edit Log'.")
    print(f"  [ ] File > Import > Upload: {BLOG_HISTORY}")
    print("  [ ] Import location: Replace current sheet. Then check the newest edit_ids are present.")
    print(f"  [ ] If meta history is also in Google Sheets, repeat for tab 'Meta Edit History' with: {META_HISTORY}")

if __name__ == "__main__":
    main()
    if "--apply" in sys.argv:
        print_sheet_reminder()

"""Generate reference/history-rows-<YYYY-MM>.json for update_edit_histories.py (Step 9).

Reads reference/remediation-briefs.json + reference/edit-lists.json (joined on blog_url)
and the two master workbooks (read-only, for the next free edit_ids). Deterministic
templates only; human judgement goes in an optional per-brief "history_note".
Never overwrites an existing output file unless --force is passed.
"""
import argparse
import json
import re
import sys
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook

from update_edit_histories import BLOG_HISTORY, META_HISTORY, find_header, find_sheet

BRIEFS_FILE = Path("reference/remediation-briefs.json")
EDIT_LISTS_FILE = Path("reference/edit-lists.json")
TITLE_FIELD = "Title 1"
DESCRIPTION_FIELD = "Meta Description 1"


def read_master(path, url_field):
    """Return [(edit_id, url)] for every record in the master's first edit_id sheet."""
    wb = load_workbook(path)
    ws = find_sheet(wb)
    header_row, cols = find_header(ws)
    records = []
    for r in range(header_row + 1, ws.max_row + 1):
        eid = ws.cell(r, cols["edit_id"]).value
        if eid not in (None, ""):
            records.append((str(eid).strip(), ws.cell(r, cols[url_field]).value or ""))
    return records


def next_seq(records, prefix):
    pattern = re.compile(rf"^{re.escape(prefix)}(\d+)$")
    used = [int(m.group(1)) for eid, _ in records if (m := pattern.match(eid))]
    return max(used, default=0) + 1


def slug(url):
    return url.rstrip("/").split("/")[-1]


def build_reason(brief, data_month):
    data = f", {data_month} data" if data_month else ""
    text = (
        f"Cannibalization (Semrush{data}): blog competes with the "
        f"{brief['commercial_page_type'].upper()} {urlparse(brief['commercial_url']).path} "
        f"(traffic risk {brief['traffic_risk']}). "
    )
    if brief["route"] == "refocus" and brief["old_focus_kw"] != brief["new_focus_kw"]:
        return text + (
            f"Focus KW '{brief['old_focus_kw']}' has commercial intent; "
            f"refocus to informational '{brief['new_focus_kw']}'."
        )
    return text + f"Focus KW '{brief['new_focus_kw']}' is informational; strengthen it."


def build_edit_summary(edits):
    n_meta = sum(1 for e in edits if "semji_field" in e)
    summary = f"{'; '.join(e['label'] for e in edits)}. {len(edits) - n_meta} body edits + {n_meta} metadata."
    if not any(e["label"] == "H1" for e in edits):
        summary += " H1 unchanged."
    return summary


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--report-month", required=True, help="YYYY-MM")
    ap.add_argument("--date", required=True, help="YYYY-MM-DD; used for every date field")
    ap.add_argument("--out", help="default: reference/history-rows-<report-month>.json")
    ap.add_argument("--data-month", help='Semrush data month named in "reason", e.g. "Aug 2026"')
    ap.add_argument("--force", action="store_true", help="overwrite an existing --out file")
    args = ap.parse_args()

    out = Path(args.out) if args.out else Path(f"reference/history-rows-{args.report_month}.json")
    if out.exists() and not args.force:
        sys.exit(f"{out} exists; pass --force to overwrite")

    briefs = json.loads(BRIEFS_FILE.read_text(encoding="utf-8"))["briefs"]
    edit_lists = {e["blog_url"]: e for e in json.loads(EDIT_LISTS_FILE.read_text(encoding="utf-8"))}

    blog_records = read_master(BLOG_HISTORY, "blog_url")
    meta_records = read_master(META_HISTORY, "page_url")
    blog_seq = next_seq(blog_records, f"{args.report_month}-")
    meta_seq = next_seq(meta_records, f"{args.report_month}-M")
    listed = {slug(url) for eid, url in blog_records if eid.startswith(f"{args.report_month}-") and url}

    blog_rows, meta_rows = [], []
    for brief in briefs:
        url = brief["blog_url"]
        if url not in edit_lists:
            sys.exit(f"No edit list for {url}")
        edit_list = edit_lists[url]
        if slug(url) in listed:
            print(f"WARNING: {slug(url)} already has a {args.report_month} row in the blog master", file=sys.stderr)

        edit_id = f"{args.report_month}-{blog_seq:03d}"
        blog_seq += 1
        notes = (
            f"cannibalization-engine Stage 4, rank {brief['rank']} of {len(briefs)}. "
            f"Route: {brief['route']}. Pair: blog vs {brief['commercial_page_type'].upper()}."
        )
        if brief.get("history_note"):
            notes += f" {brief['history_note']}"
        blog_rows.append({
            "edit_id": edit_id,
            "blog_url": url,
            "date_suggested": args.date,
            "report_month": args.report_month,
            "reason": build_reason(brief, args.data_month),
            "edit_summary": build_edit_summary(edit_list["edits"]),
            "edit_doc": brief["docx_link"],
            "semji_link": brief["semji_link"],
            "latest_version_published_semji": "",
            "pdp_links_location": "none_added",
            "status": "sent_to_client",
            "date_status_changed": args.date,
            "notes": notes,
        })

        meta_edits = {e["semji_field"]: e for e in edit_list["edits"] if "semji_field" in e}
        if meta_edits:
            title, description = meta_edits[TITLE_FIELD], meta_edits[DESCRIPTION_FIELD]
            meta_rows.append({
                "edit_id": f"{args.report_month}-M{meta_seq:03d}",
                "page_url": urlparse(url).path,
                "page_type": "blog",
                "report_month": args.report_month,
                "current_meta_title": title["current"],
                "proposed_meta_title": title["replace_with"],
                "current_meta_description": description["current"],
                "proposed_meta_description": description["replace_with"],
                "status": "proposed",
                "date_proposed": args.date,
                "notes": (
                    f"cannibalization-engine Stage 4 (edit {edit_id}). Route: {brief['route']}; "
                    + (
                        f"focus KW '{brief['old_focus_kw']}' -> '{brief['new_focus_kw']}'."
                        if brief["old_focus_kw"] != brief["new_focus_kw"]
                        else f"focus KW '{brief['new_focus_kw']}'."
                    )
                ),
            })
            meta_seq += 1

    out.write_text(
        json.dumps(
            {
                "report_month": args.report_month,
                "date": args.date,
                "blog_edit_history": blog_rows,
                "meta_edit_history": meta_rows,
            },
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    print(f"wrote {out}: {len(blog_rows)} blog rows, {len(meta_rows)} meta rows")


if __name__ == "__main__":
    main()

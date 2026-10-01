"""Generate output/content-brief-rows-<YYYY-MM>.xlsx: paste-ready Content Brief rows (columns A-R).

Reads reference/remediation-briefs.json + reference/edit-lists.json (joined on blog_url).
Row 1 is the header; select A2 to the last cell and paste into the live sheet's month section.
Never overwrites an existing output file unless --force is passed.
"""
import argparse
import json
import re
import os
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from build_history_rows import BRIEFS_FILE, DESCRIPTION_FIELD, EDIT_LISTS_FILE, TITLE_FIELD, slug

ROOT = Path(__file__).resolve().parents[1]
EDIT_HISTORY_DIR = Path(os.environ.get("EDIT_HISTORY_DIR", "data/edit-history"))
BLOG_CATEGORY_MAP = EDIT_HISTORY_DIR / "blog-category-map.json"

HEADERS = [
    "From", "Status", "Topic", "Central Campaign / Category", "1º KWs", "2º KWs", "Task",
    "URL to work on", "Meta Title proposal", "Meta Description proposal", "H1 proposal",
    "Content Structure", "Related Questions // Etiquetas Schema",
    "Internal Links to build to the main URL", "Anchor text for those links", "Notes",
    "Competitors", "SEMJI link",
]
WIDTHS = [10, 24, 32, 19, 24, 40, 24, 40, 40, 50, 40, 20, 14, 14, 14, 14, 14, 52]
LEAVE_AS_IS = "[LEAVE AS IS]"

# blog-category-map "primary" -> the live sheet's column D spelling
CATEGORY_DISPLAY = {
    "Steam irons": "STEAM IRON",
    "Steam stations": "STEAM STATION",
    "Full size steamers": "FULL SIZE STEAMER",
    "Handheld steamers": "HANDHELD STEAMER",
    "Fans": "FAN",
    "Vacuum cleaners": "VACUUM",
}

LOWER_WORDS = {"a", "an", "the", "of", "to", "and", "or", "for", "in", "at", "by", "vs"}
ACRONYMS = {"ac", "diy", "pdp", "usa"}


def title_case(text):
    words = text.split()
    out = []
    for i, w in enumerate(words):
        if w.lower() in ACRONYMS:
            out.append(w.upper())
        elif w.lower() in LOWER_WORDS and 0 < i < len(words) - 1:
            out.append(w.lower())
        else:
            out.append(w[:1].upper() + w[1:])
    return " ".join(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--report-month", required=True, help="YYYY-MM")
    ap.add_argument("--out", help="default: output/content-brief-rows-<report-month>.xlsx")
    ap.add_argument("--force", action="store_true", help="overwrite an existing --out file")
    args = ap.parse_args()

    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", args.report_month):
        sys.exit("--report-month must be YYYY-MM")
    out = Path(args.out) if args.out else Path(f"output/content-brief-rows-{args.report_month}.xlsx")
    if out.exists() and not args.force:
        sys.exit(f"{out} exists; pass --force to overwrite")

    briefs = sorted(json.loads(BRIEFS_FILE.read_text(encoding="utf-8"))["briefs"], key=lambda b: b["rank"])
    edit_lists = {e["blog_url"]: e for e in json.loads(EDIT_LISTS_FILE.read_text(encoding="utf-8"))}
    primary = {
        b["path"]: b["primary"]
        for b in json.loads(BLOG_CATEGORY_MAP.read_text(encoding="utf-8"))["blogs"]
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "Content Brief rows"
    ws.append(HEADERS)

    blank_category, blank_link = [], []
    for brief in briefs:
        url = brief["blog_url"]
        if url not in edit_lists:
            sys.exit(f"No edit list for {url}")
        edit_list = edit_lists[url]
        if slug(url) != edit_list["blog_slug"]:
            sys.exit(f"Slug mismatch for {url}: edit list says {edit_list['blog_slug']}")

        meta = {e["semji_field"]: e["replace_with"] for e in edit_list["edits"] if "semji_field" in e}
        for field in (TITLE_FIELD, DESCRIPTION_FIELD):
            if field not in meta:
                print(f"WARNING: {slug(url)} has no '{field}' edit; using {LEAVE_AS_IS}", file=sys.stderr)
        h1 = next((e["replace_with"] for e in edit_list["edits"] if e["label"] == "H1"), LEAVE_AS_IS)

        category = CATEGORY_DISPLAY.get(primary.get(f"/blog/post/{slug(url)}"), "")
        if not category:
            blank_category.append(slug(url))
        if not brief.get("semji_link"):
            blank_link.append(slug(url))

        ws.append([
            "SEO",
            "Optimization / Published",
            title_case(brief["new_focus_kw"]),
            category,
            brief["new_focus_kw"],
            "; ".join(brief["fan_outs"]),
            "Optimize an existing page",
            "/" + slug(url),
            meta.get(TITLE_FIELD, LEAVE_AS_IS),
            meta.get(DESCRIPTION_FIELD, LEAVE_AS_IS),
            h1,
            None,
            "SEE SEMJI", "SEE SEMJI", "SEE SEMJI", "SEE SEMJI",
            None,
            brief.get("semji_link") or None,
        ])

    header_fill = PatternFill("solid", fgColor="4C1E7F")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    ws.freeze_panes = "A2"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}: {len(briefs)} rows")

    print("\nHITL checklist:")
    print(f"  [ ] Paste A2 to R{len(briefs) + 1} into the {args.report_month} section of the live Content Brief.")
    if blank_category:
        print(f"  [ ] Fill blank column D for: {', '.join(blank_category)}")
    if blank_link:
        print(f"  [ ] Fill blank column R (SEMJI link) for: {', '.join(blank_link)}")
    print("  [ ] Confirm none of these pages already has a row in that month section.")


if __name__ == "__main__":
    main()

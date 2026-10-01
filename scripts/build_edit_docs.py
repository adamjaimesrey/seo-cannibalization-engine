"""Render reference/edit-lists.json to output/edit-lists/<slug>-edit-list.docx and verify quotes."""
import csv
import json
import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
EDIT_LISTS = ROOT / "reference" / "edit-lists.json"
SEMJI_CSV = ROOT / "reference" / "semji-focus-keywords.csv"
OUT_DIR = ROOT / "output" / "edit-lists"
MAX_COUNTED = 10


def norm(text):
    text = text.replace("’", "'").replace("‘", "'")
    text = text.replace("“", '"').replace("”", '"')
    text = text.replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", text).strip()


def load_semji():
    with open(SEMJI_CSV, newline="", encoding="utf-8-sig") as f:
        return {row["URL"]: row for row in csv.DictReader(f)}


def verify(entry, semji):
    source = norm((ROOT / entry["blog_copy_file"]).read_text(encoding="utf-8"))
    row = semji.get(entry["blog_url"], {})
    failures = []
    for e in entry["edits"]:
        if e["kind"] == "insert":
            continue
        if e["kind"] == "metadata":
            ok = norm(e["current"]) == norm(row.get(e["semji_field"], ""))
        else:
            ok = norm(e["current"]) in source
        if not ok:
            failures.append(e["number"])
    counted = [e for e in entry["edits"] if e["kind"] != "metadata"]
    cov = entry["coverage"]
    missing_entities = [x for x in entry["entities"] if not cov["entities"].get(x)]
    missing_fan_outs = [x for x in entry["fan_outs"] if not cov["fan_outs"].get(x)]
    return failures, len(counted), missing_entities, missing_fan_outs


def add_field(doc, label, value):
    p = doc.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run(value)


def render(entry):
    doc = Document()
    doc.add_heading(f"Edit list: {entry['blog_slug']}", level=1)
    add_field(doc, "Blog URL", entry["blog_url"])
    add_field(doc, "Commercial URL", f"{entry['commercial_url']} ({entry['commercial_page_type']})")
    add_field(doc, "Route", entry["route"])
    add_field(doc, "Focus KW", f"{entry['old_focus_kw']} -> {entry['new_focus_kw']}")
    add_field(doc, "Focus prompt", entry["focus_prompt"])
    add_field(doc, "Query fan-outs", "")
    for fo in entry["fan_outs"]:
        doc.add_paragraph(fo, style="List Bullet")
    add_field(doc, "Entities", ", ".join(entry["entities"]))

    for e in entry["edits"]:
        doc.add_heading(f"EDIT {e['number']} — {e['label']}", level=2)
        add_field(doc, "Where", e["where"])
        add_field(doc, "Current", e["current"])
        add_field(doc, "Replace with", e["replace_with"])
        add_field(doc, "Why", e["why"])

    doc.add_heading("Coverage map", level=2)
    rows = [("Entity", n, v) for n, v in entry["coverage"]["entities"].items()]
    rows += [("Fan-out", n, v) for n, v in entry["coverage"]["fan_outs"].items()]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Type", "Item", "Edit(s)")):
        cell.text = text
    for kind, name, nums in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = kind, name
        cells[2].text = ", ".join(str(n) for n in nums)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"{entry['blog_slug']}-edit-list.docx"
    doc.save(path)
    return path


def main():
    entries = json.loads(EDIT_LISTS.read_text(encoding="utf-8"))
    semji = load_semji()
    for entry in entries:
        failures, counted, miss_ent, miss_fo = verify(entry, semji)
        path = render(entry)
        print(f"[{entry['rank']}] {entry['blog_slug']} -> {path.relative_to(ROOT)}")
        print(f"  quote check: {'OK' if not failures else 'FAILED edits ' + str(failures)}")
        print(f"  edits excluding meta title/description: {counted} (max {MAX_COUNTED})"
              f"{'' if counted <= MAX_COUNTED else '  OVER CAP'}")
        print(f"  entities covered: {'yes' if not miss_ent else 'NO, missing ' + str(miss_ent)}")
        print(f"  fan-outs covered: {'yes' if not miss_fo else 'NO, missing ' + str(miss_fo)}")


if __name__ == "__main__":
    main()
